import openpyxl
import pytest

from src.validation import build_validation_report

KNOWN_MISSING_NUMBERS = {
    7, 23, 46, 67, 71, 74, 82, 83, 96, 105, 106, 107,
    131, 147, 155, 190, 192, 193, 196, 198, 202,
}


def test_known_missing_numbers_detected(default_workbook_path, gdf, matcher):
    report = build_validation_report(default_workbook_path, gdf, matcher)
    assert set(report.missing_numbers) == KNOWN_MISSING_NUMBERS


def test_missing_numbers_are_not_assigned_a_country(default_workbook_path, gdf, matcher):
    """Never fabricate a country for a missing chronological number."""
    report = build_validation_report(default_workbook_path, gdf, matcher)
    mapped_numbers = {m.number for m in report.mapped}
    assert mapped_numbers.isdisjoint(report.missing_numbers)


def test_bundled_workbook_mostly_maps_cleanly(default_workbook_path, gdf, matcher):
    report = build_validation_report(default_workbook_path, gdf, matcher)
    assert report.mapped_count >= 180
    assert report.ambiguous_count == 0


def test_duplicate_visit_numbers_detected(tmp_path, gdf, matcher):
    path = tmp_path / "dup_numbers.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Number", "Country / Territory", "Region(s)"])
    ws.append([1, "Iraq", "Middle East & Asia"])
    ws.append([1, "Syria", "Middle East & Asia"])
    ws.append([2, "Oman", "Middle East & Asia"])
    wb.save(path)

    report = build_validation_report(path, gdf, matcher)
    assert 1 in report.duplicate_numbers
    assert set(report.duplicate_numbers[1]) == {"Iraq", "Syria"}


def test_duplicate_country_names_recorded_not_discarded(tmp_path, gdf, matcher):
    path = tmp_path / "dup_countries.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Number", "Country / Territory", "Region(s)"])
    ws.append([1, "Oman", "Middle East & Asia"])
    ws.append([2, "Iraq", "Middle East & Asia"])
    ws.append([3, "Oman", "Middle East & Asia"])
    wb.save(path)

    report = build_validation_report(path, gdf, matcher)
    assert report.duplicate_country_names.get("Oman") == [1, 3]
    # Both visits are still counted individually -- not silently collapsed.
    assert sum(1 for m in report.mapped if m.country_raw == "Oman") == 2


def test_unmatched_name_reported_not_dropped(tmp_path, gdf, matcher):
    path = tmp_path / "unmatched.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Number", "Country / Territory", "Region(s)"])
    ws.append([1, "Iraq", "Middle East & Asia"])
    ws.append([2, "Freedonia", "Middle East & Asia"])
    wb.save(path)

    report = build_validation_report(path, gdf, matcher)
    assert report.unmatched_count == 1
    assert report.unresolved[0].country_raw == "Freedonia"
